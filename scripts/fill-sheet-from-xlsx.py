#!/usr/bin/env python3
"""Copy an .xlsx workbook into an EMPTY Google Sheet through the Sheets API.

    python scripts/fill-sheet-from-xlsx.py data/species-base/Spathiphyllum-base.xlsx <sheet id>
    python scripts/fill-sheet-from-xlsx.py <xlsx> <sheet id> --apply

The cloud route for a new genus book (2026-09-29). Uploading the .xlsx through the Drive
connector means passing the file as base64 in a tool call, which a session cannot do
reliably, and the service account has no Drive API. So the owner-side connector creates an
empty Google Sheet at the My Drive root and shares it to the service account as writer, and
this script fills it: every tab in order, cell values (formulas stay formulas, text stays
text), fonts, fills, alignment and wrap, column widths, row heights, frozen panes and the
basic filter, which is what Drive's own xlsx conversion carries over for these books.

Refuses a target that holds any value, so it can never overwrite a working book. Dry run
by default; --apply writes. Key: ~/.gcp/aroidpedia-sheets.json (--key to override).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

DEFAULT_KEY = Path.home() / ".gcp" / "aroidpedia-sheets.json"
H_ALIGN = {"left": "LEFT", "center": "CENTER", "right": "RIGHT"}
V_ALIGN = {"top": "TOP", "center": "MIDDLE", "bottom": "BOTTOM"}


def rgb(color) -> dict | None:
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    hexa = color.rgb[-6:]
    if color.rgb in ("00000000",):
        return None
    return {k: int(hexa[i:i + 2], 16) / 255 for k, i in (("red", 0), ("green", 2), ("blue", 4))}


def cell_data(cell) -> dict:
    out: dict = {}
    v = cell.value
    if v is not None:
        if isinstance(v, str) and v.startswith("="):
            out["userEnteredValue"] = {"formulaValue": v}
        elif isinstance(v, bool):
            out["userEnteredValue"] = {"boolValue": v}
        elif isinstance(v, (int, float)):
            out["userEnteredValue"] = {"numberValue": v}
        else:
            out["userEnteredValue"] = {"stringValue": str(v)}
    fmt: dict = {}
    f = cell.font
    tf: dict = {}
    if f is not None:
        if f.name:
            tf["fontFamily"] = f.name
        if f.sz:
            tf["fontSize"] = float(f.sz)
        if f.b:
            tf["bold"] = True
        c = rgb(f.color)
        if c:
            tf["foregroundColor"] = c
    if tf:
        fmt["textFormat"] = tf
    if cell.fill is not None and cell.fill.fill_type == "solid":
        c = rgb(cell.fill.fgColor)
        if c:
            fmt["backgroundColor"] = c
    a = cell.alignment
    if a is not None:
        if a.horizontal in H_ALIGN:
            fmt["horizontalAlignment"] = H_ALIGN[a.horizontal]
        if a.vertical in V_ALIGN:
            fmt["verticalAlignment"] = V_ALIGN[a.vertical]
        if a.wrap_text:
            fmt["wrapStrategy"] = "WRAP"
    if fmt:
        out["userEnteredFormat"] = fmt
    return out


def sheet_requests(ws, sheet_id: int) -> list:
    reqs = []
    ncols = ws.max_column
    if ws.auto_filter.ref:
        ncols = max(ncols, range_boundaries(ws.auto_filter.ref)[2])
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append({"values": [cell_data(ws.cell(r, c)) for c in range(1, ws.max_column + 1)]})
    reqs.append({"updateCells": {
        "rows": rows,
        "fields": "userEnteredValue,userEnteredFormat",
        "start": {"sheetId": sheet_id, "rowIndex": 0, "columnIndex": 0}}})
    for col, dim in ws.column_dimensions.items():
        if dim.width:
            i = openpyxl.utils.column_index_from_string(col) - 1
            reqs.append({"updateDimensionProperties": {
                "range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
                "properties": {"pixelSize": int(dim.width * 7 + 5)}, "fields": "pixelSize"}})
    for r, dim in ws.row_dimensions.items():
        if dim.height:
            reqs.append({"updateDimensionProperties": {
                "range": {"sheetId": sheet_id, "dimension": "ROWS", "startIndex": r - 1, "endIndex": r},
                "properties": {"pixelSize": round(dim.height * 4 / 3)}, "fields": "pixelSize"}})
    if ws.freeze_panes:
        col, row = range_boundaries(ws.freeze_panes + ":" + ws.freeze_panes)[:2]
        reqs.append({"updateSheetProperties": {
            "properties": {"sheetId": sheet_id,
                           "gridProperties": {"frozenRowCount": row - 1, "frozenColumnCount": col - 1}},
            "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}})
    if ws.auto_filter.ref:
        c1, r1, c2, r2 = range_boundaries(ws.auto_filter.ref)
        reqs.append({"setBasicFilter": {"filter": {"range": {
            "sheetId": sheet_id, "startRowIndex": r1 - 1, "endRowIndex": r2,
            "startColumnIndex": c1 - 1, "endColumnIndex": c2}}}})
    return reqs, ncols


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("sheet_id")
    ap.add_argument("--key", type=Path, default=DEFAULT_KEY)
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()

    import gspread
    wb = openpyxl.load_workbook(a.xlsx)
    sh = gspread.service_account(filename=str(a.key)).open_by_key(a.sheet_id)
    existing = sh.worksheets()
    held = [w.title for w in existing if any(any(v for v in row) for row in w.get_all_values())]
    if held:
        sys.exit(f"{sh.title}: tabs {held} already hold values; this script only fills an empty sheet")
    print(f"target: {sh.title} ({a.sheet_id}), tabs now {[w.title for w in existing]}")
    for ws in wb.worksheets:
        print(f"  {ws.title}: {ws.max_row} rows x {ws.max_column} cols, freeze {ws.freeze_panes}, "
              f"filter {ws.auto_filter.ref}")
    if not a.apply:
        print("dry run: nothing written (--apply to write)")
        return

    # tabs: rename the default first tab, add the rest, drop any other empty default tab
    targets = []
    for i, ws in enumerate(wb.worksheets):
        rows, cols = max(ws.max_row, 1000), max(ws.max_column, 26)
        if ws.auto_filter.ref:
            cols = max(cols, range_boundaries(ws.auto_filter.ref)[2])
        if i == 0:
            w = existing[0]
            w.update_title(ws.title)
            w.resize(rows=rows, cols=cols)
        else:
            w = sh.add_worksheet(title=ws.title, rows=rows, cols=cols)
        targets.append((ws, w))
    for w in existing[1:]:
        sh.del_worksheet(w)
    for ws, w in targets:
        reqs, _ = sheet_requests(ws, w.id)
        sh.batch_update({"requests": reqs})
        print(f"  wrote {ws.title}")
    print("tabs:", [w.title for w in sh.worksheets()])


if __name__ == "__main__":
    main()
